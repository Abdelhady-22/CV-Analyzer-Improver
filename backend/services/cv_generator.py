"""
CV Generator — builds a brand-new ATS-safe DOCX from structured data.
Single-column, Calibri, standard bullets, no tables/images/graphics.
"""

from __future__ import annotations

import uuid
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from models.cv import StructuredCVData
from config import ensure_upload_dir


# ── Styling constants ──────────────────────────────────────────

FONT_NAME = "Calibri"
FONT_SIZE_BODY = Pt(11)
FONT_SIZE_HEADING = Pt(13)
FONT_SIZE_NAME = Pt(16)
FONT_COLOR = RGBColor(0, 0, 0)
MARGIN = Inches(0.75)


def _add_styled_paragraph(
    doc: Document,
    text: str,
    font_size=FONT_SIZE_BODY,
    bold=False,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    space_after=Pt(4),
):
    """Add a paragraph with ATS-safe styling."""
    para = doc.add_paragraph()
    para.alignment = alignment
    para.paragraph_format.space_after = space_after
    para.paragraph_format.line_spacing = 1.15
    run = para.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = font_size
    run.font.color.rgb = FONT_COLOR
    run.bold = bold
    return para


def _add_section_heading(doc: Document, title: str):
    """Add a section heading with a bottom border feel."""
    _add_styled_paragraph(
        doc, title.upper(), font_size=FONT_SIZE_HEADING,
        bold=True, space_after=Pt(2),
    )
    # Add a thin separator line
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.space_before = Pt(0)
    run = para.add_run("─" * 60)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(150, 150, 150)


def _add_bullet(doc: Document, text: str):
    """Add a bullet point paragraph."""
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing = 1.15
    # Clear default run and add styled one
    if para.runs:
        para.runs[0].text = text
        para.runs[0].font.name = FONT_NAME
        para.runs[0].font.size = FONT_SIZE_BODY
        para.runs[0].font.color.rgb = FONT_COLOR
    else:
        run = para.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = FONT_SIZE_BODY
        run.font.color.rgb = FONT_COLOR


def generate_cv(cv_data: StructuredCVData) -> str:
    """
    Generate a brand-new ATS-safe DOCX from structured CV data.
    Returns the path to the generated file.
    """
    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = MARGIN
        section.bottom_margin = MARGIN
        section.left_margin = MARGIN
        section.right_margin = MARGIN

    # ── Name & Contact ─────────────────────────────────────
    _add_styled_paragraph(
        doc, cv_data.full_name,
        font_size=FONT_SIZE_NAME, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=Pt(2),
    )

    contact_parts = []
    if cv_data.email:
        contact_parts.append(cv_data.email)
    if cv_data.phone:
        contact_parts.append(cv_data.phone)
    if cv_data.location:
        contact_parts.append(cv_data.location)
    if cv_data.linkedin:
        contact_parts.append(cv_data.linkedin)
    if cv_data.website:
        contact_parts.append(cv_data.website)

    if contact_parts:
        _add_styled_paragraph(
            doc, " | ".join(contact_parts),
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            space_after=Pt(8),
        )

    # ── Summary ────────────────────────────────────────────
    if cv_data.summary:
        _add_section_heading(doc, "Professional Summary")
        _add_styled_paragraph(doc, cv_data.summary, space_after=Pt(8))

    # ── Experience ─────────────────────────────────────────
    if cv_data.experience:
        _add_section_heading(doc, "Professional Experience")
        for exp in cv_data.experience:
            # Job title + company
            title_line = f"{exp.job_title} — {exp.company}"
            _add_styled_paragraph(doc, title_line, bold=True, space_after=Pt(1))
            # Dates + location
            date_line = f"{exp.start_date} – {exp.end_date}"
            if exp.location:
                date_line += f" | {exp.location}"
            _add_styled_paragraph(doc, date_line, space_after=Pt(2))
            # Bullets
            for bullet in exp.bullets:
                _add_bullet(doc, bullet)

    # ── Education ──────────────────────────────────────────
    if cv_data.education:
        _add_section_heading(doc, "Education")
        for edu in cv_data.education:
            edu_line = f"{edu.degree} — {edu.institution}"
            _add_styled_paragraph(doc, edu_line, bold=True, space_after=Pt(1))
            date_line = f"{edu.start_date} – {edu.end_date}"
            if edu.location:
                date_line += f" | {edu.location}"
            if edu.gpa:
                date_line += f" | GPA: {edu.gpa}"
            _add_styled_paragraph(doc, date_line, space_after=Pt(4))

    # ── Skills ─────────────────────────────────────────────
    if cv_data.skills:
        _add_section_heading(doc, "Skills")
        _add_styled_paragraph(
            doc, " • ".join(cv_data.skills), space_after=Pt(8)
        )

    # ── Certifications ─────────────────────────────────────
    if cv_data.certifications:
        _add_section_heading(doc, "Certifications")
        for cert in cv_data.certifications:
            cert_line = cert.name
            if cert.issuer:
                cert_line += f" — {cert.issuer}"
            if cert.date:
                cert_line += f" ({cert.date})"
            _add_bullet(doc, cert_line)

    # Save
    output_dir = ensure_upload_dir()
    filename = f"generated_cv_{uuid.uuid4().hex[:8]}.docx"
    output_path = output_dir / filename
    doc.save(str(output_path))

    return str(output_path)
